"""
Export of phenotyping results to XLSX, PDF, CSV and JSON.

The builders are pure functions over an ExportDataset, so they can be tested
without a database and reused by every export endpoint.
"""
import csv
import io
import json
import os
from dataclasses import dataclass, field
from datetime import datetime, timezone
from typing import Dict, List, Optional, Tuple
from xml.sax.saxutils import escape

from app.config import PDF_FONT_PATH
from app.trait_analysis import PRIMARY_TRAITS, collect_records, summarize

PLATFORM_NAME = "PhenoPlant"
PDF_MAX_ROWS = 1000

BASE_COLUMNS = [
    ("study", "Эксперимент"),
    ("plant", "Растение"),
    ("treatment", "Обработка"),
    ("capturedAt", "Дата и время съёмки (UTC)"),
    ("image", "Файл изображения"),
    ("folder", "Папка"),
    ("solution", "Метод фенотипирования"),
    ("runDate", "Дата анализа"),
    ("runId", "ID прогона"),
    ("sha256", "SHA-256 изображения"),
]


@dataclass
class ExportDataset:
    studies: List[dict]
    columns: List[dict]          # value columns: {"key", "title", "unit", "coTermId"}
    rows: List[dict]             # base fields + "values": {column key: value}
    summary: dict
    generated_at: str
    generated_by: str = ""
    filters: Dict[str, str] = field(default_factory=dict)
    thumbnails: List[Tuple[str, bytes]] = field(default_factory=list)
    records: List[dict] = field(default_factory=list)  # trait records (height, leaf area) per image

    @property
    def title(self) -> str:
        names = [s.get("name", "") for s in self.studies if s.get("name")]
        return ", ".join(names) if names else "Результаты фенотипирования"


# ─── Dataset assembly ────────────────────────────────────────────

def _in_range(day: str, date_from: str, date_to: str) -> bool:
    if not day:
        return not (date_from or date_to)
    if date_from and day < date_from:
        return False
    if date_to and day > date_to:
        return False
    return True


def build_dataset(
    studies: List[dict],
    runs: List[dict],
    images: Dict[str, dict],
    plants: Dict[str, dict],
    *,
    variable_names: Optional[List[str]] = None,
    co_term_ids: Optional[List[str]] = None,
    date_from: str = "",
    date_to: str = "",
    generated_by: str = "",
) -> ExportDataset:
    names = set(variable_names or [])
    co_terms = set(co_term_ids or [])
    filter_vars = bool(names or co_terms)
    study_names = {s["_id"]: s.get("name", s["_id"]) for s in studies}

    columns: Dict[str, dict] = {}
    rows: List[dict] = []
    kept_runs: List[dict] = []
    for run in sorted(runs, key=lambda r: r.get("createdAt", "")):
        kept_results = []
        for res in run.get("results", []):
            image = images.get(res.get("imageId", ""), {})
            captured = image.get("capturedAt", "")
            day = (captured or run.get("createdAt", ""))[:10]
            if not _in_range(day, date_from, date_to):
                continue
            plant = plants.get(image.get("observationUnitId", ""), {})
            values = {}
            for m in res.get("measurements", []):
                name = m.get("variableName", "")
                if not name:
                    continue
                if filter_vars and name not in names and m.get("coTermId", "") not in co_terms:
                    continue
                if name not in columns:
                    columns[name] = {"key": name, "title": name, "unit": m.get("unit", ""), "coTermId": m.get("coTermId", "")}
                values[name] = m.get("value", "")
            if filter_vars and not values:
                continue
            kept_results.append(res)
            rows.append({
                "study": study_names.get(run.get("studyId", ""), run.get("studyId", "")),
                "plant": plant.get("name") or res.get("plantName", ""),
                "treatment": plant.get("treatment", ""),
                "capturedAt": captured,
                "image": res.get("fileName", ""),
                "folder": run.get("folderName", ""),
                "solution": run.get("solutionName", ""),
                "runDate": run.get("createdAt", ""),
                "runId": run.get("_id") or run.get("id", ""),
                "sha256": image.get("sha256", ""),
                "values": values,
                "imageId": res.get("imageId", ""),
                "annotatedImageId": res.get("annotatedImageId", ""),
            })
        if kept_results:
            kept_runs.append({**run, "results": kept_results})

    # CropOntology variables first, then raw algorithm metrics
    ordered = sorted(columns.values(), key=lambda c: (0 if c["coTermId"] else 1))
    for col in ordered:
        if col["unit"]:
            col["title"] = f"{col['key']}, {col['unit']}"

    records = collect_records(kept_runs, images, plants)
    summary = summarize(records)
    filters = {}
    if filter_vars:
        filters["Переменные"] = ", ".join(sorted(names | co_terms))
    if date_from or date_to:
        filters["Период"] = f"{date_from or '…'} — {date_to or '…'}"
    return ExportDataset(
        studies=studies,
        columns=ordered,
        rows=rows,
        summary=summary,
        generated_at=datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
        generated_by=generated_by,
        filters=filters,
        records=records,
    )


# ─── CSV / JSON ──────────────────────────────────────────────────

def build_csv(ds: ExportDataset) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow([key for key, _ in BASE_COLUMNS] + [c["key"] for c in ds.columns])
    for row in ds.rows:
        writer.writerow([row.get(key, "") for key, _ in BASE_COLUMNS] + [row["values"].get(c["key"], "") for c in ds.columns])
    return out.getvalue().encode("utf-8-sig")  # BOM so Excel opens Cyrillic correctly


def build_json(ds: ExportDataset) -> bytes:
    by_key = {c["key"]: c for c in ds.columns}
    observations = []
    for row in ds.rows:
        for key, value in row["values"].items():
            col = by_key[key]
            observations.append({
                "studyName": row["study"],
                "observationUnitName": row["plant"],
                "treatment": row["treatment"],
                "observationVariableName": key,
                "observationVariablePUI": col["coTermId"],
                "unit": col["unit"],
                "value": value,
                "observationTimeStamp": row["capturedAt"],
                "imageFileName": row["image"],
                "imageSha256": row["sha256"],
                "analysisRunId": row["runId"],
                "phenotypingMethod": row["solution"],
            })
    payload = {
        "metadata": {
            "platform": PLATFORM_NAME,
            "generatedAt": ds.generated_at,
            "generatedBy": ds.generated_by,
            "studies": [{"studyDbId": s["_id"], "studyName": s.get("name", "")} for s in ds.studies],
            "filters": ds.filters,
        },
        "observations": observations,
        "summary": ds.summary,
    }
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")


# ─── XLSX ────────────────────────────────────────────────────────

def build_xlsx(ds: ExportDataset) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F6B3F")
    title_font = Font(bold=True, size=13)

    def style_header(ws, row: int, ncols: int):
        for col in range(1, ncols + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def p_format(ws, col: int):
        cell = ws.cell(row=ws.max_row, column=col)
        if isinstance(cell.value, float):
            cell.number_format = "0.00E+00" if 0 < cell.value < 1e-4 else "0.0000"

    def print_setup(ws):
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.gridLines = False

    def autosize(ws, min_width=8, max_width=45):
        widths: Dict[int, int] = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    widths[cell.column] = max(widths.get(cell.column, 0), len(str(cell.value)))
        for col, width in widths.items():
            ws.column_dimensions[get_column_letter(col)].width = max(min_width, min(width + 2, max_width))

    wb = Workbook()

    # Results
    ws = wb.active
    ws.title = "Результаты"
    headers = [title for _, title in BASE_COLUMNS] + [c["title"] for c in ds.columns]
    ws.append(headers)
    style_header(ws, 1, len(headers))
    for row in ds.rows:
        values = []
        for c in ds.columns:
            v = row["values"].get(c["key"], None)
            values.append(v if isinstance(v, (int, float)) else (v if v not in ("", None) else None))
        ws.append([row.get(key, "") for key, _ in BASE_COLUMNS] + values)
    first_value_col = len(BASE_COLUMNS) + 1
    for r in range(2, ws.max_row + 1):
        for c in range(first_value_col, len(headers) + 1):
            if isinstance(ws.cell(row=r, column=c).value, float):
                ws.cell(row=r, column=c).number_format = "0.00"
    ws.freeze_panes = "B2"
    if ds.rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    autosize(ws)

    # Statistics
    st = wb.create_sheet("Статистика")
    st.append(["Описательная статистика по фенотипическим признакам"])
    st["A1"].font = title_font
    stat_headers = ["Признак", "Ед.", "Группа", "n", "Среднее", "SD", "SEM", "Мин", "Медиана", "Макс", "CV, %", "95% ДИ (нижн.)", "95% ДИ (верхн.)"]
    st.append([])
    st.append(stat_headers)
    style_header(st, 3, len(stat_headers))

    def stat_row(trait, group, d):
        return [trait["name"], trait["unit"], group, d["n"], d["mean"], d["sd"], d["sem"], d["min"], d["median"],
                d["max"], d["cv"], d["ci95Low"], d["ci95High"]]

    for trait in ds.summary["traits"]:
        st.append(stat_row(trait, "Все растения", trait["overall"]))
        for g in trait["byTreatment"]:
            st.append(stat_row(trait, g["treatment"], g))

    comparisons = [(t, c) for t in ds.summary["traits"] for c in t["comparisons"]]
    st.append([])
    st.append(["Сравнение групп с контролем (t-критерий Уэлча, α = 0,05)"])
    st.cell(row=st.max_row, column=1).font = title_font
    if comparisons:
        cmp_headers = ["Признак", "Группа", "Контроль", "Разность средних", "t", "df", "p", "Различие значимо"]
        st.append(cmp_headers)
        style_header(st, st.max_row, len(cmp_headers))
        for t, c in comparisons:
            st.append([t["name"], c["group"], c["reference"], c["meanDiff"], c["t"], c["df"], c["pValue"],
                       "да" if c["significant"] else "нет"])
            p_format(st, 7)
    else:
        st.append(["Недостаточно данных: нужны минимум две группы обработки по 2 и более растения."])

    by_date = [(t, c) for t in ds.summary["traits"] for c in t.get("comparisonsByDate", [])]
    if by_date:
        st.append([])
        st.append(["Сравнение с контролем внутри каждой даты съёмки (t-критерий Уэлча, α = 0,05)"])
        st.cell(row=st.max_row, column=1).font = title_font
        bd_headers = ["Признак", "Дата съёмки", "Группа", "Контроль", "Разность средних", "t", "df", "p", "Различие значимо"]
        st.append(bd_headers)
        style_header(st, st.max_row, len(bd_headers))
        for t, c in by_date:
            st.append([t["name"], c["date"], c["group"], c["reference"], c["meanDiff"], c["t"], c["df"], c["pValue"],
                       "да" if c["significant"] else "нет"])
            p_format(st, 8)

    corr = ds.summary.get("correlation")
    st.append([])
    st.append(["Корреляция высоты растения и площади проекции листьев (Пирсон)"])
    st.cell(row=st.max_row, column=1).font = title_font
    if corr:
        st.append(["n", "r", "p", "Связь значима"])
        style_header(st, st.max_row, 4)
        st.append([corr["n"], corr["r"], corr["pValue"], "да" if corr["significant"] else "нет"])
        p_format(st, 3)
    else:
        st.append(["Недостаточно парных измерений (нужно не менее 3)."])
    autosize(st, max_width=40)

    # Dynamics with native Excel charts
    dyn = wb.create_sheet("Динамика")
    dates = ds.summary.get("dates", [])
    dyn_headers = ["Дата съёмки"]
    for t in ds.summary["traits"]:
        dyn_headers += [f"{t['name']}, среднее ({t['unit']})", f"{t['name']}, SD", f"{t['name']}, n"]
    dyn.append(dyn_headers)
    style_header(dyn, 1, len(dyn_headers))
    dyn_index = {t["key"]: {d["date"]: d for d in t["dynamics"]} for t in ds.summary["traits"]}
    for day in dates:
        row = [day]
        for t in ds.summary["traits"]:
            d = dyn_index[t["key"]].get(day)
            row += [d["mean"], d["sd"], d["n"]] if d else [None, None, None]
        dyn.append(row)
    autosize(dyn, max_width=28)
    if len(dates) >= 2:
        for i, t in enumerate(ds.summary["traits"]):
            chart = LineChart()
            chart.title = f"{t['name']}: среднее по дате съёмки"
            chart.y_axis.title = t["unit"]
            chart.x_axis.title = "Дата"
            chart.height, chart.width = 7.5, 16
            col = 2 + i * 3
            chart.add_data(Reference(dyn, min_col=col, min_row=1, max_row=len(dates) + 1), titles_from_data=True)
            chart.set_categories(Reference(dyn, min_col=1, min_row=2, max_row=len(dates) + 1))
            dyn.add_chart(chart, f"{get_column_letter(len(dyn_headers) + 2)}{2 + i * 16}")

    # Growth per plant
    gr = wb.create_sheet("Рост растений")
    gr_headers = ["Признак", "Растение", "Обработка", "Первая дата", "Последняя дата", "Дней",
                  "Начальное значение", "Конечное значение", "Абс. скорость роста, ед./сут", "Отн. скорость роста, 1/сут"]
    gr.append(gr_headers)
    style_header(gr, 1, len(gr_headers))
    for t in ds.summary["traits"]:
        for g in t["growth"]:
            gr.append([t["name"], g["plant"], g["treatment"], g["firstDate"], g["lastDate"], g["days"],
                       g["firstValue"], g["lastValue"], g["agr"], g["rgr"]])
    autosize(gr, max_width=30)

    # Export metadata
    meta = wb.create_sheet("Сведения")
    meta.append(["Сведения о выгрузке"])
    meta["A1"].font = title_font
    info = [
        ("Платформа", PLATFORM_NAME),
        ("Сформировано", ds.generated_at),
        ("Пользователь", ds.generated_by),
        ("Эксперименты", ds.title),
        ("Строк результатов", len(ds.rows)),
        ("Изображений", ds.summary.get("images", 0)),
        ("Растений", ds.summary.get("plants", 0)),
        ("Дат съёмки", len(dates)),
    ] + list(ds.filters.items()) + [
        ("Контроль целостности", "Для каждого изображения указан SHA-256, вычисленный при загрузке"),
    ]
    for k, v in info:
        meta.append([k, v])
        meta.cell(row=meta.max_row, column=1).font = Font(bold=True)
    meta.column_dimensions["A"].width = 24
    meta.column_dimensions["B"].width = 80

    for sheet in wb.worksheets:
        print_setup(sheet)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ─── PDF ─────────────────────────────────────────────────────────

_FONT_CANDIDATES = [
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("/usr/share/fonts/TTF/DejaVuSans.ttf", "/usr/share/fonts/TTF/DejaVuSans-Bold.ttf"),
    ("/usr/share/fonts/dejavu/DejaVuSans.ttf", "/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf"),
    ("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf", "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"),
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
]
_fonts_registered: Optional[Tuple[str, str]] = None


def _register_fonts() -> Tuple[str, str]:
    """Register a TTF font with Cyrillic glyphs; the built-in PDF fonts have none."""
    global _fonts_registered
    if _fonts_registered:
        return _fonts_registered
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = list(_FONT_CANDIDATES)
    if PDF_FONT_PATH:
        bold = PDF_FONT_PATH.replace(".ttf", "-Bold.ttf")
        candidates.insert(0, (PDF_FONT_PATH, bold if os.path.exists(bold) else PDF_FONT_PATH))
    for regular, bold in candidates:
        if os.path.exists(regular):
            pdfmetrics.registerFont(TTFont("PP-Regular", regular))
            pdfmetrics.registerFont(TTFont("PP-Bold", bold if os.path.exists(bold) else regular))
            _fonts_registered = ("PP-Regular", "PP-Bold")
            return _fonts_registered
    raise RuntimeError("No TTF font with Cyrillic glyphs found; set PDF_FONT_PATH or install fonts-dejavu-core")


def _fmt(v, digits: int = 2) -> str:
    if v is None or v == "":
        return "—"
    if isinstance(v, float):
        return f"{v:.{digits}f}".replace(".", ",")
    return str(v)


def _fmt_p(p) -> str:
    if p is None:
        return "—"
    return "< 0,0001" if p < 0.0001 else _fmt(float(p), 4)


def _p_clause(p) -> str:
    text = _fmt_p(p)
    return f"p {text}" if text.startswith("<") else f"p = {text}"


def _thumbnail(data: bytes, max_side: int = 700) -> Optional[bytes]:
    from PIL import Image
    try:
        with Image.open(io.BytesIO(data)) as img:
            img = img.convert("RGB")
            img.thumbnail((max_side, max_side))
            out = io.BytesIO()
            img.save(out, "JPEG", quality=80)
            return out.getvalue()
    except Exception:
        return None


def build_pdf(ds: ExportDataset) -> bytes:
    from reportlab.graphics.charts.lineplots import LinePlot
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.charts.legends import Legend
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.graphics.widgets.markers import makeMarker
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (Image as RLImage, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    regular, bold = _register_fonts()
    green = colors.HexColor("#2F6B3F")
    light = colors.HexColor("#EEF5EF")
    palette = [colors.HexColor("#2F6B3F"), colors.HexColor("#C8811A"), colors.HexColor("#3A6EA5"),
               colors.HexColor("#A23B72"), colors.HexColor("#5B8C5A")]

    s_title = ParagraphStyle("t", fontName=bold, fontSize=16, leading=20, spaceAfter=4)
    s_sub = ParagraphStyle("s", fontName=regular, fontSize=10, leading=13, textColor=colors.HexColor("#444444"))
    s_h = ParagraphStyle("h", fontName=bold, fontSize=12.5, leading=16, spaceBefore=10, spaceAfter=5, textColor=green)
    s_p = ParagraphStyle("p", fontName=regular, fontSize=9.5, leading=12.5)
    s_cell = ParagraphStyle("c", fontName=regular, fontSize=8, leading=10)
    s_cell_b = ParagraphStyle("cb", fontName=bold, fontSize=8, leading=10, textColor=colors.white)
    s_cap = ParagraphStyle("cap", fontName=regular, fontSize=8, leading=10, alignment=TA_CENTER)

    def P(text, style=s_cell):
        return Paragraph(escape(str(text)), style)

    def table(data, widths, header_rows=1):
        body = [[P(c, s_cell_b) for c in data[0]]] + [[P(c) for c in r] for r in data[1:]]
        t = Table(body, colWidths=widths, repeatRows=header_rows)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), green),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, light]),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B8C4BB")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
        ]))
        return t

    width = A4[0] - 30 * mm
    story = [
        Paragraph("Отчёт о результатах цифрового фенотипирования", s_title),
        Paragraph(escape(ds.title), s_sub),
        Spacer(1, 6),
    ]

    dates = ds.summary.get("dates", [])
    meta_rows = [["Параметр", "Значение"],
                 ["Эксперименты", ds.title],
                 ["Период съёмки", f"{dates[0]} — {dates[-1]}" if dates else "—"],
                 ["Изображений / растений", f"{ds.summary.get('images', 0)} / {ds.summary.get('plants', 0)}"],
                 ["Строк результатов", str(len(ds.rows))],
                 ["Сформировано", f"{ds.generated_at}{(', ' + ds.generated_by) if ds.generated_by else ''}"]]
    meta_rows += [[k, v] for k, v in ds.filters.items()]
    story.append(table(meta_rows, [55 * mm, width - 55 * mm]))

    # 1. Descriptive statistics
    story.append(Paragraph("1. Описательная статистика признаков", s_h))
    stat = [["Признак", "Группа", "n", "Среднее", "SD", "Мин", "Медиана", "Макс", "CV, %", "95% ДИ"]]
    for t in ds.summary["traits"]:
        for group, d in [("Все растения", t["overall"])] + [(g["treatment"], g) for g in t["byTreatment"]]:
            ci = f"{_fmt(d['ci95Low'])} … {_fmt(d['ci95High'])}" if d["ci95Low"] is not None else "—"
            stat.append([f"{t['name']}, {t['unit']}", group, d["n"], _fmt(d["mean"]), _fmt(d["sd"]), _fmt(d["min"]),
                         _fmt(d["median"]), _fmt(d["max"]), _fmt(d["cv"], 1), ci])
    story.append(table(stat, [34 * mm, 22 * mm, 8 * mm, 16 * mm, 14 * mm, 14 * mm, 16 * mm, 16 * mm, 12 * mm, width - 152 * mm]))

    comparisons = [(t, c) for t in ds.summary["traits"] for c in t["comparisons"]]
    if comparisons:
        story.append(Paragraph("2. Сравнение групп обработки с контролем (t-критерий Уэлча)", s_h))
        cmp_rows = [["Признак", "Группа", "Контроль", "Разность средних", "t", "df", "p", "Значимо при α=0,05"]]
        for t, c in comparisons:
            cmp_rows.append([t["name"], c["group"], c["reference"], _fmt(c["meanDiff"]), _fmt(c["t"]), _fmt(c["df"], 1),
                             _fmt_p(c["pValue"]), "да" if c["significant"] else "нет"])
        story.append(table(cmp_rows, [35 * mm, 25 * mm, 25 * mm, 22 * mm, 14 * mm, 13 * mm, 19 * mm, width - 153 * mm]))
    by_date = [(t, c) for t in ds.summary["traits"] for c in t.get("comparisonsByDate", [])]
    if by_date:
        story.append(Spacer(1, 6))
        story.append(Paragraph("Сравнение выше выполнено по всем датам съёмки вместе. Ниже то же сравнение внутри каждой даты, "
                               "для растений одного возраста.", s_p))
        story.append(Spacer(1, 4))
        rows = [["Признак", "Дата съёмки", "Группа", "Разность средних", "t", "df", "p", "Значимо"]]
        for t, c in by_date:
            rows.append([t["name"], c["date"], c["group"], _fmt(c["meanDiff"]), _fmt(c["t"]), _fmt(c["df"], 1),
                         _fmt_p(c["pValue"]), "да" if c["significant"] else "нет"])
        story.append(table(rows, [35 * mm, 22 * mm, 28 * mm, 22 * mm, 14 * mm, 13 * mm, 19 * mm, width - 153 * mm]))

    # 3. Charts
    def line_chart(trait) -> Drawing:
        all_dates = [d["date"] for d in trait["dynamics"] if d["mean"] is not None]
        index = {day: i for i, day in enumerate(all_dates)}
        series = [(g["treatment"], [(index[p["date"]], p["mean"]) for p in g["points"] if p["mean"] is not None])
                  for g in trait.get("dynamicsByTreatment", [])]
        series = [(name, pts) for name, pts in series if pts]
        if len(series) <= 1:
            series = [("Все растения", [(i, d["mean"]) for i, d in enumerate(d_ for d_ in trait["dynamics"] if d_["mean"] is not None)])]
        h = 64 * mm
        d = Drawing(width, h)
        d.add(String(0, h - 10, f"{trait['name']}, {trait['unit']} — среднее по дате съёмки", fontName=bold, fontSize=9))
        lp = LinePlot()
        lp.x, lp.y, lp.width, lp.height = 30, 22, width - 150, h - 45
        lp.data = [pts for _, pts in series]
        for i in range(len(series)):
            color = palette[i % len(palette)]
            lp.lines[i].strokeColor = color
            lp.lines[i].strokeWidth = 1.6
            lp.lines[i].symbol = makeMarker("FilledCircle")
            lp.lines[i].symbol.fillColor = color
            lp.lines[i].symbol.size = 4
        lp.xValueAxis.valueMin, lp.xValueAxis.valueMax = -0.3, max(len(all_dates) - 0.7, 0.7)
        lp.xValueAxis.valueSteps = list(range(len(all_dates)))
        lp.xValueAxis.labelTextFormat = lambda v: all_dates[int(round(v))][5:] if 0 <= int(round(v)) < len(all_dates) else ""
        for ax in (lp.xValueAxis, lp.yValueAxis):
            ax.labels.fontName = regular
            ax.labels.fontSize = 7
        lp.yValueAxis.valueMin = 0
        d.add(lp)
        legend = Legend()
        legend.x, legend.y = width - 110, h - 30
        legend.fontName, legend.fontSize = regular, 7.5
        legend.alignment = "right"
        legend.columnMaximum = 8
        legend.colorNamePairs = [(palette[i % len(palette)], name[:24]) for i, (name, _) in enumerate(series)]
        d.add(legend)
        return d

    def bar_chart(trait) -> Drawing:
        groups = [g for g in trait["byTreatment"] if g["mean"] is not None]
        d = Drawing(width, 58 * mm)
        d.add(String(0, 58 * mm - 10, f"{trait['name']}, {trait['unit']} — среднее по группам", fontName=bold, fontSize=9))
        bc = VerticalBarChart()
        bc.x, bc.y, bc.width, bc.height = 30, 22, width - 50, 58 * mm - 45
        bc.data = [[g["mean"] for g in groups]]
        bc.categoryAxis.categoryNames = [g["treatment"][:18] for g in groups]
        bc.categoryAxis.labels.fontName = regular
        bc.categoryAxis.labels.fontSize = 7
        bc.valueAxis.labels.fontName = regular
        bc.valueAxis.labels.fontSize = 7
        bc.valueAxis.valueMin = 0
        bc.bars[0].fillColor = green
        bc.barWidth = 10
        d.add(bc)
        return d

    charts = []
    for t in ds.summary["traits"]:
        if len([x for x in t["dynamics"] if x["mean"] is not None]) >= 2:
            charts.append(line_chart(t))
        elif t["byTreatment"] and any(g["mean"] is not None for g in t["byTreatment"]):
            charts.append(bar_chart(t))
    corr = ds.summary.get("correlation")
    if charts or corr:
        story.append(Paragraph("3. Графики", s_h))
        for c in charts:
            story.append(KeepTogether([c, Spacer(1, 4)]))
    if corr:
        by_image: Dict[str, Dict[str, float]] = {}
        for rec in ds.records:
            by_image.setdefault(rec["imageId"], {})[rec["trait"]] = rec["value"]
        pts = [(v["height"], v["leaf_area"]) for v in by_image.values() if "height" in v and "leaf_area" in v]
        if len(pts) >= 3:
            d = Drawing(width, 62 * mm)
            d.add(String(0, 62 * mm - 10, f"Высота (см) и площадь проекции листьев (см²): r = {_fmt(corr['r'], 3)}, "
                                          f"{_p_clause(corr['pValue'])}, n = {corr['n']}", fontName=bold, fontSize=9))
            sp = LinePlot()
            sp.x, sp.y, sp.width, sp.height = 30, 22, width - 50, 62 * mm - 45
            sp.data = [pts]
            sp.joinedLines = 0
            sp.lines[0].symbol = makeMarker("FilledCircle")
            sp.lines[0].symbol.fillColor = palette[2]
            sp.lines[0].symbol.size = 3.5
            sp.lines[0].strokeColor = None
            for ax in (sp.xValueAxis, sp.yValueAxis):
                ax.labels.fontName = regular
                ax.labels.fontSize = 7
                ax.valueMin = 0
            d.add(sp)
            story.append(KeepTogether([d, Spacer(1, 4)]))
        else:
            story.append(Paragraph(f"Корреляция высоты и площади: r = {_fmt(corr['r'], 3)}, {_p_clause(corr['pValue'])}, n = {corr['n']}.", s_p))

    # 4. Annotated images
    thumbs = [(cap, _thumbnail(data)) for cap, data in ds.thumbnails]
    thumbs = [(cap, data) for cap, data in thumbs if data]
    if thumbs:
        story.append(Paragraph("4. Примеры аннотированных изображений", s_h))
        cells, row = [], []
        cell_w = (width - 6 * mm) / 2
        for cap, data in thumbs:
            img = RLImage(io.BytesIO(data))
            ratio = img.imageHeight / float(img.imageWidth)
            img.drawWidth = cell_w - 4 * mm
            img.drawHeight = min(img.drawWidth * ratio, 58 * mm)
            img.drawWidth = img.drawHeight / ratio
            row.append([img, Paragraph(escape(cap), s_cap)])
            if len(row) == 2:
                cells.append(row)
                row = []
        if row:
            cells.append(row + [""])
        grid = Table(cells, colWidths=[cell_w, cell_w])
        grid.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("ALIGN", (0, 0), (-1, -1), "CENTER")]))
        story.append(grid)

    # 5. Measurements
    story.append(PageBreak())
    story.append(Paragraph("5. Результаты измерений", s_h))
    trait_values: Dict[tuple, float] = {(rec["imageId"], rec["trait"]): rec["value"] for rec in ds.records}
    head = ["№", "Растение", "Обработка", "Дата съёмки", "Файл"] + [f"{t['name']}, {t['unit']}" for t in PRIMARY_TRAITS]
    rows = [head]
    for i, r in enumerate(ds.rows[:PDF_MAX_ROWS], 1):
        rows.append([i, r["plant"] or "—", r["treatment"] or "—", (r["capturedAt"] or "—")[:16].replace("T", " "),
                     r["image"]] + [_fmt(trait_values.get((r["imageId"], t["key"]))) for t in PRIMARY_TRAITS])
    value_w = 26 * mm
    fixed = 10 * mm + 28 * mm + 26 * mm + 26 * mm + value_w * len(PRIMARY_TRAITS)
    story.append(table(rows, [10 * mm, 28 * mm, 26 * mm, 26 * mm, max(width - fixed, 25 * mm)] + [value_w] * len(PRIMARY_TRAITS)))
    if len(ds.rows) > PDF_MAX_ROWS:
        story.append(Spacer(1, 4))
        story.append(Paragraph(f"Показаны первые {PDF_MAX_ROWS} строк из {len(ds.rows)}; полная таблица — в выгрузке XLSX.", s_p))

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont(regular, 7.5)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawString(15 * mm, 9 * mm, f"{PLATFORM_NAME} · платформа цифрового фенотипирования · {ds.generated_at}")
        canvas.drawRightString(A4[0] - 15 * mm, 9 * mm, f"стр. {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=14 * mm,
                            bottomMargin=16 * mm, title="Отчёт о результатах фенотипирования", author=PLATFORM_NAME)
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()


MEDIA_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
    "csv": "text/csv; charset=utf-8",
    "json": "application/json",
}

BUILDERS = {"xlsx": build_xlsx, "pdf": build_pdf, "csv": build_csv, "json": build_json}
