import csv
import hashlib
import io
import json
import zipfile

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app.export import build_csv, build_dataset, build_json, build_pdf, build_xlsx
from tests.test_analysis import seed_experiment


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


async def analysed_study(client, headers, db):
    study = await seed_experiment(client, headers, db)
    resp = await client.post(f"/api/studies/{study['id']}/run-analysis",
                             json={"solution_id": "sol-rgb", "folder_name": "Съёмка"}, headers=headers)
    assert resp.status_code == 200
    return study


def pdf_text(content: bytes) -> str:
    return "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(content)).pages)


def small_dataset():
    studies = [{"_id": "s1", "name": "Опыт"}]
    images = {f"i{i}": {"capturedAt": f"2026-03-0{1 + (i % 2) * 7}T09:00:00Z", "observationUnitId": f"ou{i % 2}",
                        "sha256": "ab" * 32} for i in range(4)}
    plants = {"ou0": {"name": "P1", "treatment": "Контроль"}, "ou1": {"name": "P2", "treatment": "Засуха"}}
    results = [{"imageId": f"i{i}", "fileName": f"img{i}.png", "plantName": "",
                "metrics": {"height_cm": 10.0 + i, "leaf_area_cm2": 40.0 + 3 * i},
                "measurements": [{"coTermId": "CO_321:0000020", "variableName": "Plant height", "value": 10.0 + i, "unit": "см"},
                                 {"coTermId": "CO_321:0000080", "variableName": "Leaf area", "value": 40.0 + 3 * i, "unit": "см²"},
                                 {"coTermId": "", "variableName": "green_index", "value": 120.5, "unit": ""}]}
               for i in range(4)]
    runs = [{"_id": "r1", "studyId": "s1", "createdAt": "2026-03-09 10:00", "solutionName": "RGB", "folderName": "F",
             "results": results}]
    return studies, runs, images, plants


class TestBuilders:
    def test_dataset_columns_and_filters(self):
        ds = build_dataset(*small_dataset())
        assert [c["key"] for c in ds.columns] == ["Plant height", "Leaf area", "green_index"]
        assert ds.columns[0]["title"] == "Plant height, см"
        assert len(ds.rows) == 4 and ds.rows[0]["plant"] == "P1"

        only_area = build_dataset(*small_dataset(), co_term_ids=["CO_321:0000080"])
        assert [c["key"] for c in only_area.columns] == ["Leaf area"]

        march_1 = build_dataset(*small_dataset(), date_from="2026-03-01", date_to="2026-03-01")
        assert len(march_1.rows) == 2
        assert march_1.filters["Период"] == "2026-03-01 — 2026-03-01"

    def test_xlsx_structure(self):
        wb = load_workbook(io.BytesIO(build_xlsx(build_dataset(*small_dataset()))))
        assert wb.sheetnames == ["Результаты", "Статистика", "Динамика", "Рост растений", "Сведения"]
        ws = wb["Результаты"]
        header = [c.value for c in ws[1]]
        assert header[:2] == ["Эксперимент", "Растение"] and "Plant height, см" in header
        assert ws.max_row == 5
        assert ws.freeze_panes == "B2" and ws.auto_filter.ref
        height_cell = ws.cell(row=2, column=header.index("Plant height, см") + 1)
        assert height_cell.value == 10.0 and height_cell.number_format == "0.00"
        stats_text = " ".join(str(c.value) for row in wb["Статистика"].iter_rows() for c in row if c.value)
        assert "Высота растения" in stats_text and "Площадь проекции листовой поверхности" in stats_text
        assert "Уэлча" in stats_text
        assert len(wb["Динамика"]._charts) == 2

    def test_pdf_contains_cyrillic_report(self):
        content = build_pdf(build_dataset(*small_dataset(), generated_by="admin@test.ru"))
        assert content.startswith(b"%PDF")
        text = pdf_text(content)
        assert "Отчёт о результатах цифрового фенотипирования" in text
        assert "Высота растения" in text and "Контроль" in text
        assert "Результаты измерений" in text

    def test_csv_and_json(self):
        ds = build_dataset(*small_dataset())
        raw = build_csv(ds)
        assert raw.startswith("﻿".encode())
        rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
        assert rows[0][:3] == ["study", "plant", "treatment"] and len(rows) == 5
        payload = json.loads(build_json(ds))
        assert payload["metadata"]["studies"][0]["studyName"] == "Опыт"
        assert len(payload["observations"]) == 12
        assert payload["observations"][0]["observationVariablePUI"] == "CO_321:0000020"


class TestEndpoints:
    async def test_study_export_formats(self, client, auth_headers, db):
        study = await analysed_study(client, auth_headers, db)
        xlsx = await client.get(f"/api/studies/{study['id']}/export?format=xlsx", headers=auth_headers)
        assert xlsx.status_code == 200 and xlsx.headers["content-type"] == XLSX
        assert "attachment;" in xlsx.headers["content-disposition"]
        assert load_workbook(io.BytesIO(xlsx.content))["Результаты"].max_row == 9

        pdf = await client.get(f"/api/studies/{study['id']}/export?format=pdf", headers=auth_headers)
        assert pdf.status_code == 200 and pdf.headers["content-type"] == "application/pdf"
        text = pdf_text(pdf.content)
        assert "Соя, засуха" in text
        assert "Примеры аннотированных изображений" in text
        assert "Сравнение групп" in text

        bad = await client.get(f"/api/studies/{study['id']}/export?format=docx", headers=auth_headers)
        assert bad.status_code == 422

    async def test_post_export_with_filters(self, client, auth_headers, db):
        study = await analysed_study(client, auth_headers, db)
        resp = await client.post("/api/export", json={
            "format": "csv", "study_ids": [study["id"]], "variable_names": ["Plant height"],
            "date_from": "2026-03-08", "date_to": "2026-03-31",
        }, headers=auth_headers)
        assert resp.status_code == 200
        rows = list(csv.reader(io.StringIO(resp.content.decode("utf-8-sig"))))
        assert len(rows) == 5 and rows[0][-1] == "Plant height"

        by_var_id = await db.client["phenoplant_test_brapi"].observation_variables.insert_one(
            {"_id": "v-area", "name": "Leaf area", "coTermId": "CO_321:0000080"})
        assert by_var_id
        resp = await client.post("/api/export", json={"format": "json", "study_ids": [study["id"]],
                                                      "variable_ids": ["v-area"]}, headers=auth_headers)
        names = {o["observationVariableName"] for o in resp.json()["observations"]}
        assert names == {"Leaf area"}

    async def test_post_export_validation(self, client, auth_headers):
        assert (await client.post("/api/export", json={"format": "xlsx", "study_ids": []}, headers=auth_headers)).status_code == 422
        assert (await client.post("/api/export", json={"format": "xlsx", "study_ids": ["nope"]}, headers=auth_headers)).status_code == 404
        assert (await client.post("/api/export", json={"format": "xlsx", "study_ids": ["x"], "date_from": "01.03.2026"},
                                  headers=auth_headers)).status_code == 422
        assert (await client.post("/api/export", json={"format": "xlsx", "study_ids": ["x"]})).status_code == 403

    async def test_zip_package_with_manifest(self, client, auth_headers, db):
        study = await analysed_study(client, auth_headers, db)
        resp = await client.get(f"/api/studies/{study['id']}/export-zip", headers=auth_headers)
        zf = zipfile.ZipFile(io.BytesIO(resp.content))
        names = set(zf.namelist())
        assert {"results.csv", "results.xlsx", "report.pdf", "manifest.json"} <= names
        assert len([n for n in names if n.startswith("dataset/")]) == 8
        assert len([n for n in names if n.startswith("results/")]) == 8
        manifest = json.loads(zf.read("manifest.json"))
        for path, digest in manifest["files"].items():
            assert hashlib.sha256(zf.read(path)).hexdigest() == digest

    async def test_legacy_csv_endpoint(self, client, auth_headers, db):
        study = await analysed_study(client, auth_headers, db)
        resp = await client.get(f"/api/studies/{study['id']}/export-results-csv", headers=auth_headers)
        assert resp.status_code == 200 and resp.content.startswith("﻿".encode())
